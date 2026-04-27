import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def build_resume(data, output_path="output/resume.docx"):
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.3)
        section.bottom_margin = Inches(0.3)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)

    def set_font(run, size, bold=False, color=None):
        run.font.name = "Calibri"
        run.font.size = Pt(size)
        run.font.bold = bold
        if color:
            run.font.color.rgb = RGBColor(*color)

    def clean_text(value):
        return str(value).strip() if value is not None else ""

    def format_gpa(gpa_value):
        gpa_text = clean_text(gpa_value)
        if not gpa_text:
            return ""
        if gpa_text.lower().startswith("gpa"):
            return gpa_text
        try:
            num = float(gpa_text)
            if 0 <= num <= 4.0:
                return f"GPA: {num:.1f}/4.0"
        except ValueError:
            pass
        return f"GPA: {gpa_text}"

    def add_section_header(doc, title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(1)
        run = p.add_run(title.upper())
        set_font(run, 10, bold=True)
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '6')
        bottom.set(qn('w:space'), '1')
        bottom.set(qn('w:color'), '000000')
        pBdr.append(bottom)
        pPr.append(pBdr)
        return p

    def add_bullet(doc, text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Inches(0.2)
        run = p.add_run(f"\u2022  {text}")
        set_font(run, 9)
        return p

    def add_hyperlink(paragraph, text, url):
        part = paragraph.part
        r_id = part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
        hyperlink = OxmlElement('w:hyperlink')
        hyperlink.set(qn('r:id'), r_id)
        run = OxmlElement('w:r')
        rPr = OxmlElement('w:rPr')
        rStyle = OxmlElement('w:rStyle')
        rStyle.set(qn('w:val'), 'Hyperlink')
        rPr.append(rStyle)
        run.append(rPr)
        t = OxmlElement('w:t')
        t.text = text
        run.append(t)
        hyperlink.append(run)
        paragraph._p.append(hyperlink)

    # ── NAME ──
    name_p = doc.add_paragraph()
    name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_p.paragraph_format.space_after = Pt(2)
    name_run = name_p.add_run(data['name'])
    set_font(name_run, 14, bold=True)

    # ── CONTACT INFO ──
    contact_p = doc.add_paragraph()
    contact_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact_p.paragraph_format.space_after = Pt(4)
    contact_parts = []
    for key in ("phone", "email", "location"):
        value = clean_text(data.get(key))
        if value:
            contact_parts.append(value)
    if contact_parts:
        r1 = contact_p.add_run(" | ".join(contact_parts))
        set_font(r1, 10)

    has_linkedin = bool(clean_text(data.get("linkedin")))
    has_github = bool(clean_text(data.get("github")))
    if has_linkedin or has_github:
        if contact_parts:
            sep = contact_p.add_run(" | ")
            set_font(sep, 10)
        if has_linkedin:
            add_hyperlink(contact_p, "LinkedIn", data["linkedin"])
        if has_github:
            if has_linkedin:
                mid = contact_p.add_run(" | ")
                set_font(mid, 10)
            add_hyperlink(contact_p, "GitHub", data["github"])

    # ── PROFESSIONAL SUMMARY ──
    add_section_header(doc, "Professional Summary")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    run = p.add_run(data['summary'])
    set_font(run, 9)

    # ── TECHNICAL SKILLS ──
    add_section_header(doc, "Technical Skills")
    for category, items in data['skills'].items():
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        bold_run = p.add_run(f"{category}: ")
        set_font(bold_run, 10, bold=True)
        normal_run = p.add_run(items)
        set_font(normal_run, 10)

    # ── WORK EXPERIENCE ──
    add_section_header(doc, "Work Experience")
    for exp in data['experience']:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(0)
        left_run = p.add_run(f"{exp['company']} | {exp['location']}")
        set_font(left_run, 10, bold=True)
        tab_run = p.add_run(f"  {exp['dates']}")
        set_font(tab_run, 10, bold=True)
        p2 = doc.add_paragraph()
        p2.paragraph_format.space_before = Pt(0)
        p2.paragraph_format.space_after = Pt(0)
        title_run = p2.add_run(exp['title'])
        set_font(title_run, 10, bold=True)
        for bullet in exp['bullets']:
            add_bullet(doc, bullet)

    # ── ACADEMIC PROJECTS ──
    add_section_header(doc, "Academic Projects")
    for proj in data['projects']:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(0)
        proj_run = p.add_run(proj['name'])
        set_font(proj_run, 10, bold=True)
        for bullet in proj['bullets']:
            add_bullet(doc, bullet)

    # ── AWARDS & ACTIVITIES ──
    add_section_header(doc, "Awards & Activities")
    for award in data['awards']:
        add_bullet(doc, award)

    # ── EDUCATION ──
    add_section_header(doc, "Academic Qualification")
    for edu in data['education']:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(0)
        degree_run = p.add_run(edu['degree'])
        set_font(degree_run, 10, bold=True)
        date_run = p.add_run(f"  {edu['dates']}")
        set_font(date_run, 10, bold=True)
        p2 = doc.add_paragraph()
        p2.paragraph_format.space_before = Pt(0)
        p2.paragraph_format.space_after = Pt(0)
        school_run = p2.add_run(edu['school'])
        set_font(school_run, 10)
        gpa_value = format_gpa(edu.get("gpa"))
        if gpa_value:
            gpa_run = p2.add_run(f"  |  {gpa_value}")
            set_font(gpa_run, 10)

    os.makedirs(os.path.dirname(output_path) if os.path.dirname(output_path) else "output", exist_ok=True)
    doc.save(output_path)
    return output_path